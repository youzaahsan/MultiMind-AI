import io
import tempfile
from pathlib import Path
import docx
import pandas as pd
from PIL import Image
from pypdf import PdfWriter

from app.multimodal.pdf_processor import PDFProcessor
from app.multimodal.docx_processor import DocxProcessor
from app.multimodal.csv_processor import CSVProcessor
from app.multimodal.excel_processor import ExcelProcessor
from app.multimodal.image_processor import ImageProcessor

def test_pdf_processor(tmp_path):
    # Generate simple PDF
    pdf_path = tmp_path / "test.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with open(pdf_path, "wb") as f:
        writer.write(f)

    proc = PDFProcessor(pdf_path)
    res = proc.extract()
    assert res["total_pages"] == 1
    assert "metadata" in res
    assert res["metadata"]["filename"] == "test.pdf"

def test_docx_processor(tmp_path):
    docx_path = tmp_path / "sample.docx"
    doc = docx.Document()
    doc.add_heading("Chapter 1: Deep Learning", level=1)
    doc.add_paragraph("Deep learning models utilize multi-layered artificial neural networks.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Model"
    table.cell(0, 1).text = "Accuracy"
    table.cell(1, 0).text = "ResNet"
    table.cell(1, 1).text = "94.5%"
    doc.save(str(docx_path))

    proc = DocxProcessor(docx_path)
    res = proc.extract()
    assert len(res["headings"]) == 1
    assert "Deep Learning" in res["headings"][0]["text"]
    assert len(res["tables"]) == 1
    assert "ResNet" in res["tables"][0]["markdown"]

def test_csv_processor(tmp_path):
    csv_path = tmp_path / "metrics.csv"
    df = pd.DataFrame({
        "epoch": [1, 2, 3, 4],
        "loss": [0.85, 0.45, 0.22, 0.15],
        "optimizer": ["Adam", "Adam", "SGD", "Adam"],
    })
    df.to_csv(csv_path, index=False)

    proc = CSVProcessor(csv_path)
    res = proc.extract()
    assert res["metadata"]["rows"] == 4
    assert res["metadata"]["columns_count"] == 3
    assert "loss" in res["statistics"]
    assert res["statistics"]["loss"]["min"] == 0.15
    assert "Adam" in res["categorical_summary"]["optimizer"]["top_values"]

def test_excel_processor(tmp_path):
    xlsx_path = tmp_path / "sales.xlsx"
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        df1 = pd.DataFrame({"product": ["A", "B"], "qty": [10, 20]})
        df2 = pd.DataFrame({"region": ["North", "South"], "quota": [100, 200]})
        df1.to_excel(writer, sheet_name="Q1", index=False)
        df2.to_excel(writer, sheet_name="Q2", index=False)

    proc = ExcelProcessor(xlsx_path)
    res = proc.extract()
    assert res["metadata"]["total_sheets"] == 2
    assert "Q1" in res["sheets"]
    assert "Q2" in res["sheets"]

def test_image_processor(tmp_path):
    img_path = tmp_path / "chart.png"
    img = Image.new("RGB", (120, 80), color="blue")
    img.save(img_path)

    proc = ImageProcessor(img_path)
    res = proc.extract()
    assert res["metadata"]["width"] == 120
    assert res["metadata"]["height"] == 80
    assert res["metadata"]["format"] == "PNG"
    assert res["base64_data"].startswith("data:image/png;base64,")
