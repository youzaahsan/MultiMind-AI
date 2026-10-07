from pathlib import Path
from typing import Any, Dict, List
import docx
from app.utils.logging import get_logger

logger = get_logger("docx_processor")

class DocxProcessor:
    """Extracts paragraphs, headings, tables, and structure from DOCX documents."""

    def __init__(self, file_path: str | Path):
        self.file_path = Path(file_path)

    def extract(self) -> Dict[str, Any]:
        if not self.file_path.exists():
            raise FileNotFoundError(f"DOCX file not found: {self.file_path}")

        doc = docx.Document(str(self.file_path))

        paragraphs: List[Dict[str, Any]] = []
        headings: List[Dict[str, Any]] = []
        tables: List[Dict[str, Any]] = []
        full_text_chunks: List[str] = []

        # Process paragraphs & headings
        for p_idx, p in enumerate(doc.paragraphs):
            text = p.text.strip()
            if not text:
                continue

            style_name = p.style.name if p.style else ""
            is_heading = "Heading" in style_name or style_name.startswith("Title")

            item = {
                "index": p_idx,
                "text": text,
                "style": style_name,
                "is_heading": is_heading,
            }

            if is_heading:
                headings.append(item)
            else:
                paragraphs.append(item)

            full_text_chunks.append(text)

        # Process tables
        for t_idx, table in enumerate(doc.tables):
            table_rows = []
            for row in table.rows:
                row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                table_rows.append(row_cells)

            if table_rows:
                # Format as markdown table
                header = table_rows[0]
                separator = ["---"] * len(header)
                md_table_lines = [
                    "| " + " | ".join(header) + " |",
                    "| " + " | ".join(separator) + " |",
                ]
                for r in table_rows[1:]:
                    md_table_lines.append("| " + " | ".join(r) + " |")

                md_table = "\n".join(md_table_lines)
                tables.append({
                    "table_index": t_idx,
                    "rows_count": len(table_rows),
                    "cols_count": len(header),
                    "markdown": md_table,
                    "raw_rows": table_rows,
                })
                full_text_chunks.append(f"\n[Table {t_idx + 1}]\n{md_table}\n")

        full_text = "\n\n".join(full_text_chunks)

        return {
            "metadata": {
                "filename": self.file_path.name,
                "file_size": self.file_path.stat().st_size,
                "paragraph_count": len(paragraphs),
                "heading_count": len(headings),
                "table_count": len(tables),
            },
            "headings": headings,
            "paragraphs": paragraphs,
            "tables": tables,
            "full_text": full_text,
            "total_words": len(full_text.split()),
        }
