from pathlib import Path
from typing import Any, Dict, List, Optional
import openpyxl
import pandas as pd
from app.utils.logging import get_logger

logger = get_logger("excel_processor")

class ExcelProcessor:
    """Processes multi-sheet Excel workbooks (.xlsx, .xls), formulas, and tables."""

    def __init__(self, file_path: str | Path):
        self.file_path = Path(file_path)

    def extract(self) -> Dict[str, Any]:
        if not self.file_path.exists():
            raise FileNotFoundError(f"Excel file not found: {self.file_path}")

        # Read sheet names and dataframes
        excel_file = pd.ExcelFile(self.file_path)
        sheet_names = excel_file.sheet_names

        # Read formulas using openpyxl data_only=False if xlsx
        formulas_found: Dict[str, List[str]] = {}
        if self.file_path.suffix.lower() == ".xlsx":
            try:
                wb = openpyxl.load_workbook(str(self.file_path), data_only=False)
                for sheet in wb.sheetnames:
                    ws = wb[sheet]
                    sheet_formulas = []
                    for row in ws.iter_rows(values_only=False):
                        for cell in row:
                            if cell.value and isinstance(cell.value, str) and cell.value.startswith("="):
                                sheet_formulas.append(f"{cell.coordinate}: {cell.value}")
                    if sheet_formulas:
                        formulas_found[sheet] = sheet_formulas[:20]  # keep top 20 formulas
            except Exception as e:
                logger.warning(f"Could not inspect formulas with openpyxl: {e}")

        sheets_data: Dict[str, Any] = {}
        full_text_parts: List[str] = [f"Excel Workbook: {self.file_path.name}\nTotal Sheets: {len(sheet_names)} ({', '.join(sheet_names)})"]

        for sheet in sheet_names:
            df = excel_file.parse(sheet)
            rows, cols = df.shape
            columns = list(df.columns)

            numeric_df = df.select_dtypes(include=["number"])
            stats_dict = {}
            if not numeric_df.empty:
                desc = numeric_df.describe().to_dict()
                for col, metrics in desc.items():
                    stats_dict[col] = {k: round(v, 4) if isinstance(v, (int, float)) else v for k, v in metrics.items()}

            preview_records = df.head(5).to_dict(orient="records")

            sheet_summary = [
                f"\n--- Sheet: '{sheet}' ---",
                f"Dimensions: {rows} rows x {cols} columns",
                f"Columns: {', '.join(str(c) for c in columns)}",
            ]
            for col, m in stats_dict.items():
                sheet_summary.append(f"Stat '{col}': mean={m.get('mean')}, min={m.get('min')}, max={m.get('max')}")

            if sheet in formulas_found:
                sheet_summary.append(f"Detected Formulas in '{sheet}': {', '.join(formulas_found[sheet][:5])}")

            full_text_parts.append("\n".join(sheet_summary))

            sheets_data[sheet] = {
                "rows": rows,
                "columns_count": cols,
                "columns": [str(c) for c in columns],
                "statistics": stats_dict,
                "sample_rows": preview_records,
                "formulas": formulas_found.get(sheet, []),
            }

        return {
            "metadata": {
                "filename": self.file_path.name,
                "file_size": self.file_path.stat().st_size,
                "sheet_names": sheet_names,
                "total_sheets": len(sheet_names),
            },
            "sheets": sheets_data,
            "full_text": "\n\n".join(full_text_parts),
        }
